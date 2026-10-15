#!/usr/bin/env python3
"""Link each O*NET career to one ISCO-08 unit group (ROADMAP 6.3, D4 = B).

Source: O*NET's "ESCO to O*NET-SOC 2019" crosswalk (data/onet/crosswalks/ESCO_to_ONET-SOC.xlsx, from
scripts/download_sources.py; built by O*NET "combining AI techniques with human validation"). Each row
links an ESCO code to an O*NET career. Most codes are ESCO occupations ("2512.3": the first 4 digits are
the ISCO-08 unit group), but O*NET also lists bare ISCO codes ("2512"): its own assignment of the career
to an ISCO group. The links are many-to-many (a career's ESCO matches span a median of 5 ISCO groups),
so one group is chosen per career by this rule, the first that applies:

  isco_row         exactly one 4-digit ISCO row: O*NET's own assignment                       (900 careers)
  isco_rows+esco   several 4-digit ISCO rows: the one most of the career's ESCO matches share
  minor_group+esco only a 3-digit ISCO row (minor group): the most common ESCO unit group inside it
  esco_majority    no ISCO row: the most common unit group among the ESCO matches
  none             a tie, or nothing inside the minor group: no group rather than a guess

Every unit group the career's matches touch is kept in `all_isco08` (code:count), because an ISCO
group is broader than an O*NET career and the app has to say so.

    python scripts/onet/build_onet_isco08.py   -> data/processed/onet_isco08.csv
"""

from __future__ import annotations

import csv
import sys
from collections import Counter, defaultdict
from pathlib import Path

import duckdb
from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(PROJECT_ROOT))

from career_rag.config import ONET_DUCKDB_PATH  # noqa: E402

CROSSWALK = PROJECT_ROOT / "data" / "onet" / "crosswalks" / "ESCO_to_ONET-SOC.xlsx"
OUTPUT = PROJECT_ROOT / "data" / "processed" / "onet_isco08.csv"
HEADER = ("ESCO/ISCO Code", "ESCO/ISCO Title", "O*NET-SOC 2019 Code", "O*NET-SOC 2019 Title")


def read_crosswalk() -> list[tuple[str, str]]:
    """(ESCO or ISCO code, O*NET-SOC code) pairs."""
    rows = list(load_workbook(CROSSWALK, read_only=True).worksheets[0].iter_rows(values_only=True))
    start = rows.index(HEADER) + 1  # the sheet starts with a title block
    return [(str(code).strip(), str(onet).strip()) for code, _, onet, _ in rows[start:] if code and onet]


def unique_top(counts: Counter) -> str | None:
    """The most common key, or None when it is tied (or there is none)."""
    ranked = counts.most_common(2)
    if not ranked or (len(ranked) == 2 and ranked[0][1] == ranked[1][1]):
        return None
    return ranked[0][0]


def choose(isco_rows: list[str], esco_groups: Counter) -> tuple[str | None, str]:
    four = sorted({code for code in isco_rows if len(code) == 4})
    three = sorted({code for code in isco_rows if len(code) == 3})
    if len(four) == 1:
        return four[0], "isco_row"
    if four:
        return unique_top(Counter({code: esco_groups[code] for code in four})), "isco_rows+esco"
    if three:
        inside = Counter({code: n for code, n in esco_groups.items() if code[:3] in three})
        return unique_top(inside), "minor_group+esco"
    return unique_top(esco_groups), "esco_majority"


def main() -> int:
    isco_rows: dict[str, list[str]] = defaultdict(list)
    esco_groups: dict[str, Counter] = defaultdict(Counter)
    for code, onet in read_crosswalk():
        if "." in code:
            esco_groups[onet][code.split(".")[0].zfill(4)] += 1
        else:
            isco_rows[onet].append(code)

    con = duckdb.connect(str(ONET_DUCKDB_PATH), read_only=True)
    careers = [code for (code,) in con.execute(
        "SELECT DISTINCT onetsoc_code FROM task_statements ORDER BY onetsoc_code").fetchall()]
    con.close()

    rows, methods = [], Counter()
    for career in careers:
        isco, method = choose(isco_rows[career], esco_groups[career])
        if isco is None:
            method = "none"
        methods[method] += 1
        touched = esco_groups[career] + Counter(code for code in isco_rows[career] if len(code) == 4)
        rows.append((career, isco or "", method, ";".join(f"{code}:{n}" for code, n in sorted(touched.items()))))

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT.open("w", encoding="utf-8", newline="") as out:
        writer = csv.writer(out)
        writer.writerow(["onetsoc_code", "isco08_code", "method", "all_isco08"])
        writer.writerows(rows)
    print(f"wrote {OUTPUT.relative_to(PROJECT_ROOT)}: {len(rows)} careers; {dict(methods)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
